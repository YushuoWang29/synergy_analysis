"""
MEGA BALANCE SEARCH — Find PC1/PC2 variance balance close to Pisa/IIT (70-85% / 10-25%).
"""
import os, sys, glob, json
import numpy as np
import scipy.io as sio
from sklearn.decomposition import PCA
from scipy.interpolate import interp1d
from scipy.linalg import orthogonal_procrustes
import openpyxl
from datetime import datetime

SAVE_DIR = 'PCA_results/experiments'
os.makedirs(SAVE_DIR, exist_ok=True)

N_JOINTS = 15
JOINT_NAMES = ["ThumbAbd","ThumbMCP","ThumbIP","ThumbOpp","IdxMCP","IdxPIP","MidMCP","MidPIP","IdxMidAbd","RingMCP","RingPIP","MidRingAbd","LitMCP","LitPIP","RingLitAbd"]
PISA_S1 = np.ones(N_JOINTS) / np.linalg.norm(np.ones(N_JOINTS))
PISA_S2_RAW = np.array([2,2,2,2, 1,1, 0,0,0, -1,-1,-1, -2,-2,-2], dtype=float)
PISA_S2 = PISA_S2_RAW / np.linalg.norm(PISA_S2_RAW)
FLEXION_IDS = [1, 2, 4, 5, 6, 7, 9, 10, 12, 13]

def load_task_info():
    wb = openpyxl.load_workbook('MyoKi Dataset/Task_categorization.xlsx')
    ws = wb.active
    categories, task_names = {}, {}
    action_cols = {'cylindrical': 11, 'spherical': 12, 'hook': 13, 'tripod': 14, 'pinch': 15, 'lumbrical': 16, 'complex': 17, 'none_action': 18}
    for row in ws.iter_rows(min_row=4, values_only=False):
        tid = row[0].value
        if tid is None: continue
        tid_int = int(tid)
        task_names[tid_int] = str(row[1].value).strip() if row[1].value else f"Task {tid_int}"
        actions = []
        for act_name, col_idx in action_cols.items():
            if row[col_idx].value is not None and str(row[col_idx].value).strip() == 'x':
                actions.append(act_name)
        categories[tid_int] = actions
    return categories, task_names

def make_pisa_target(jids, n_total=None):
    """Create Pisa targets for any joint count.
    For 15 joints: standard thumb-to-little gradient.
    For 20 joints (Ninapro): extended gradient.
    For fewer: subsets."""
    n = len(jids)
    s1 = np.ones(n) / np.linalg.norm(np.ones(n))
    # Create a smooth gradient based on joint indices
    # The standard 15-joint pattern is: 2,2,2,2, 1,1, 0,0,0, -1,-1,-1, -2,-2,-2
    if n_total is None or n_total == 15:
        raw = np.array([2,2,2,2, 1,1, 0,0,0, -1,-1,-1, -2,-2,-2])
    elif n_total == 20:
        # Ninapro 20 joints: approximate gradient
        raw = np.array([2,2,2,2, 2,1,1, 1,1, 0,0, -1,-1,-1, -2,-2,-2,-2,-2,-2])
    else:
        # Generic gradient based on index
        raw = np.linspace(2, -2, n_total)
    
    # Map jids indices to raw indices (approximate)
    s2 = raw[jids] / (np.linalg.norm(raw[jids]) + 1e-10)
    return s1, s2


categories, task_names = load_task_info()
power_types = {'cylindrical', 'spherical', 'hook'}
precision_types = {'tripod', 'pinch', 'lumbrical'}
grasp_types = power_types | precision_types

POWER_IDS = sorted([tid for tid, acts in categories.items() if any(a in power_types for a in acts) and 'none_action' not in acts])
ALL_GRASP_IDS = sorted([tid for tid, acts in categories.items() if any(a in grasp_types for a in acts) and 'none_action' not in acts])
PRECISION_IDS = sorted([tid for tid, acts in categories.items() if any(a in precision_types for a in acts) and 'none_action' not in acts and not any(a in power_types for a in acts)])
NON_GRASP_IDS = sorted([tid for tid, acts in categories.items() if 'none_action' in acts])
COMPLEX_IDS = sorted([tid for tid, acts in categories.items() if 'complex' in acts])
ALL_TASKS = sorted(categories.keys())

def get_pickup_ids():
    wb = openpyxl.load_workbook('MyoKi Dataset/Task_categorization.xlsx')
    ws = wb.active
    ids = set()
    for row in ws.iter_rows(min_row=4, values_only=False):
        tid = row[0].value; name = str(row[1].value).strip() if row[1].value else ''
        if tid is not None and 'place it on' in name.lower() and 'back' not in name.lower():
            ids.add(int(tid))
    return sorted(ids)

PICKUP_IDS = get_pickup_ids()

print("=" * 60)
print(f"POWER: {len(POWER_IDS)}, PRECISION: {len(PRECISION_IDS)}, ALL_GRASP: {len(ALL_GRASP_IDS)}")
print(f"NON_GRASP: {len(NON_GRASP_IDS)}, PICKUP: {len(PICKUP_IDS)}")

def load_myoki_trials(files, task_ids, resample_n=100):
    all_trials = []
    for fpath in files:
        mat = sio.loadmat(fpath)
        glove = np.nan_to_num(mat['glove_calibrated'])
        tasks = mat['task'].flatten().astype(int)
        hand_data = glove[:, :15]
        for tid in task_ids:
            mask = (tasks == tid)
            diff = np.diff(mask.astype(int), prepend=0, append=0)
            starts = np.where(diff > 0)[0]; ends = np.where(diff < 0)[0]
            for s, e in zip(starts, ends):
                trial = hand_data[s:e, :]
                if len(trial) >= 5:
                    x_old = np.linspace(0, 1, len(trial))
                    x_new = np.linspace(0, 1, resample_n)
                    f = interp1d(x_old, trial, axis=0, kind='linear', bounds_error=False, fill_value='extrapolate')
                    all_trials.append((f(x_new), tid, os.path.basename(fpath).replace('.mat', '')))
    return all_trials

def get_participant_files():
    pf = sorted(glob.glob('MyoKi Dataset/raw/P*.mat'))
    valid = [f'P{p:02d}' for p in [1,2,3,4,5,6,7,8,9,11,12]]
    return [f for f in pf if any(p in os.path.basename(f) for p in valid)]

FILES = get_participant_files()
print(f"Participants: {len(FILES)}")

def extract_static(trials, n_last=5):
    return [(np.mean(t[-n_last:], axis=0), tid, subj) for t, tid, subj in trials]
def extract_all_frames(trials):
    return [(t, tid, subj) for t, tid, subj in trials]
def extract_trial_means(trials):
    return [(np.mean(t, axis=0), tid, subj) for t, tid, subj in trials]
def extract_trial_std(trials):
    return [(np.std(t, axis=0), tid, subj) for t, tid, subj in trials]
def extract_trial_range(trials):
    return [(np.max(t, axis=0) - np.min(t, axis=0), tid, subj) for t, tid, subj in trials]
def extract_first_frame(trials):
    return [(t[0], tid, subj) for t, tid, subj in trials]
def extract_last_frame(trials):
    return [(t[-1], tid, subj) for t, tid, subj in trials]

def zscore_normalize(data):
    mean = np.mean(data, axis=0); std = np.std(data, axis=0)
    std[std < 1e-6] = 1.0; return (data - mean) / std
def range_normalize(data):
    mins = np.min(data, axis=0); maxs = np.max(data, axis=0)
    r = maxs - mins; r[r < 1e-6] = 1.0; return (data - mins) / r
def robust_normalize(data):
    med = np.median(data, axis=0)
    q75, q25 = np.percentile(data, [75, 25], axis=0)
    iqr = q75 - q25; iqr[iqr < 1e-6] = 1.0; return (data - med) / iqr
def mean_normalize(data):
    return data - np.mean(data, axis=0)
def none_normalize(data):
    return data

NORM_FUNCS = {'none': none_normalize, 'zscore': zscore_normalize, 'range': range_normalize, 'robust': robust_normalize, 'mean': mean_normalize}

def evaluate_pca(data_matrix, label, jids=None, norm='none', pisa_s1=None, pisa_s2=None):
    data = np.array(data_matrix, dtype=float)
    if len(data) < 5 or data.ndim < 2: return None
    if jids is not None: 
        # Ensure jids are within bounds
        max_jid = max(jids)
        if max_jid >= data.shape[1]:
            # Fall back to all available joints
            jids = list(range(data.shape[1]))
    else: 
        jids = list(range(data.shape[1]))
    n_j = data.shape[1]
    if jids is not None and max(jids) < data.shape[1]:
        data = data[:, jids]
        n_j = len(jids)
    if norm in NORM_FUNCS: data = NORM_FUNCS[norm](data)
    if np.isnan(data).any() or np.isinf(data).any(): return None
    
    n_comp = min(n_j, data.shape[0], 15)
    if n_comp < 2: return None  # Need at least 2 components
    pca = PCA(n_components=n_comp)
    pca.fit(data)
    evr = pca.explained_variance_ratio_
    cum = np.cumsum(evr)
    pc1 = pca.components_[0]; pc2 = pca.components_[1] if n_comp > 1 else np.zeros(n_j)
    
    if pisa_s1 is None or pisa_s2 is None:
        s1, s2 = make_pisa_target(np.array(jids), n_total=data_matrix.shape[1])
    else:
        s1, s2 = pisa_s1, pisa_s2

    
    PC = np.column_stack([pc1, pc2]); target = np.column_stack([s1, s2])
    R, _ = orthogonal_procrustes(PC, target); rotated = PC @ R
    
    s1_sim = float(np.dot(rotated[:, 0], s1) / (np.linalg.norm(rotated[:, 0]) * np.linalg.norm(s1) + 1e-10))
    s2_sim = float(np.dot(rotated[:, 1], s2) / (np.linalg.norm(rotated[:, 1]) * np.linalg.norm(s2) + 1e-10))
    
    scores = pca.transform(data)[:, :2] @ R.T
    rvar = np.var(scores, axis=0); rv_ratio = rvar / (rvar.sum() + 1e-10)
    total_pc12 = evr[:2].sum(); rv_scaled = rv_ratio * total_pc12
    
    pc1_n = pc1 / (np.linalg.norm(pc1) + 1e-10); pc2_n = pc2 / (np.linalg.norm(pc2) + 1e-10)
    pc1_s1_sim = float(np.dot(pc1_n, s1)); pc2_s2_sim = float(np.dot(pc2_n, s2))
    
    pc2_var_val = float(evr[1]) if len(evr) > 1 else 0
    bal_score = s1_sim + s2_sim + min(1.0, pc2_var_val * 8)
    
    result = {
        'label': label, 'n_samples': data.shape[0], 'n_joints': n_j, 'norm': norm,
        'jids': [int(j) for j in jids],
        'pc1_var': float(evr[0]), 'pc2_var': pc2_var_val,
        'pc1_pc2_sum': float(total_pc12),
        'pc1_s1_sim': pc1_s1_sim, 'pc2_s2_sim': pc2_s2_sim,
        'rot_s1_sim': s1_sim, 'rot_s2_sim': s2_sim, 'rot_sum': s1_sim + s2_sim,
        'rot_pc1_var': float(rv_scaled[0]), 'rot_pc2_var': float(rv_scaled[1]),
        'balance_score': bal_score,
        'pc1': pc1.tolist(), 'pc2': pc2.tolist(),
        'rot_s1': rotated[:, 0].tolist(), 'rot_s2': rotated[:, 1].tolist(),
        'n_90': int(np.searchsorted(cum, 0.90) + 1),
    }
    return result

def run_exp1():
    print(f"\n{'='*60}\nEXP 1: TASK × DATA REP × NORM\n{'='*60}")
    results = []
    task_sets = {'power_only': POWER_IDS, 'precision_only': PRECISION_IDS, 'power_precision': sorted(set(POWER_IDS)|set(PRECISION_IDS)),
                 'all_grasp': ALL_GRASP_IDS, 'all_tasks': sorted(set(ALL_TASKS)-{0}), 'pickup_only': PICKUP_IDS,
                 'all_grasp_nongrasp': sorted(set(ALL_GRASP_IDS)|set(NON_GRASP_IDS)), 'complex_grasp': sorted(set(COMPLEX_IDS)|set(ALL_GRASP_IDS))}
    reps = [('trialmean', extract_trial_means, 'none'), ('trialmean', extract_trial_means, 'zscore'), ('trialmean', extract_trial_means, 'range'), ('trialmean', extract_trial_means, 'robust'),
            ('static5', lambda t: extract_static(t,5), 'none'), ('static5', lambda t: extract_static(t,5), 'zscore'),
            ('static10', lambda t: extract_static(t,10), 'none'),
            ('trialsd', extract_trial_std, 'none'), ('trialsd', extract_trial_std, 'zscore'),
            ('trialrange', extract_trial_range, 'none'), ('trialrange', extract_trial_range, 'zscore'),
            ('allframes', extract_all_frames, 'none'), ('allframes', extract_all_frames, 'zscore'), ('allframes', extract_all_frames, 'range'),
            ('firstframe', extract_first_frame, 'none'), ('lastframe', extract_last_frame, 'none')]
    joint_sets = {'all15': None, 'flex10': FLEXION_IDS, 'no_abd': [0,1,2,3,4,5,6,7,9,10,12,13]}
    count = 0
    for task_name, task_ids in task_sets.items():
        print(f"\n  {task_name} ({len(task_ids)} tasks)...", end='', flush=True)
        trials = load_myoki_trials(FILES, task_ids)
        print(f" {len(trials)} trials")
        if len(trials) < 3: continue
        for rep_name, rep_fn, norm in reps:
            processed = rep_fn(trials)
            data_array = np.vstack([p[0] for p in processed])
            for js_name, jids in joint_sets.items():
                label = f"e1_{task_name}_{rep_name}_{js_name}_{norm}"; count += 1
                result = evaluate_pca(data_array, label, jids=jids, norm='none' if norm=='none' else norm)
                if result is None: continue
                if result['pc2_var'] > 0.05 or result['rot_sum'] > 1.5:
                    print(f"  [{count}] {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                results.append(result)
    return results

def run_exp2():
    print(f"\n{'='*60}\nEXP 2: TASK-AVERAGED (BALANCED)\n{'='*60}")
    results = []
    task_sets = {'power_only': POWER_IDS, 'precision_only': PRECISION_IDS, 'power_precision': sorted(set(POWER_IDS)|set(PRECISION_IDS)),
                 'all_grasp': ALL_GRASP_IDS, 'all_tasks': sorted(set(ALL_TASKS)-{0}), 'pickup_only': PICKUP_IDS,
                 'power_precision_complex': sorted(set(POWER_IDS)|set(PRECISION_IDS)|set(COMPLEX_IDS))}
    reps = [('trialmean', extract_trial_means, 'none'), ('trialmean', extract_trial_means, 'zscore'),
            ('static5', lambda t: extract_static(t,5), 'none'), ('static10', lambda t: extract_static(t,10), 'none'),
            ('trialsd', extract_trial_std, 'none'), ('trialsd', extract_trial_std, 'zscore')]
    joint_sets = {'all15': None, 'flex10': FLEXION_IDS, 'no_abd': [0,1,2,3,4,5,6,7,9,10,12,13]}
    count = 0
    for task_name, task_ids in task_sets.items():
        print(f"\n  {task_name} ({len(task_ids)} tasks)...", end='', flush=True)
        trials = load_myoki_trials(FILES, task_ids)
        print(f" {len(trials)} trials")
        if len(trials) < 3: continue
        for rep_name, rep_fn, norm in reps:
            processed = rep_fn(trials)
            # Build subject-task structure: subject -> task -> list of trial vectors
            by_subj_task = {}
            for p in processed:
                vec, tid, subj = p
                if subj not in by_subj_task: by_subj_task[subj] = {}
                if tid not in by_subj_task[subj]: by_subj_task[subj][tid] = []
                by_subj_task[subj][tid].append(vec)
            
            # Compute per-subject, per-task mean (each subject contributes one sample per task)
            # Then flatten: one sample = (subject, task) pair
            task_samples = []
            for subj, task_dict in by_subj_task.items():
                for tid, vecs in task_dict.items():
                    task_samples.append(np.mean(vecs, axis=0))
            data_array = np.array(task_samples)
            
            for js_name, jids in joint_sets.items():
                label = f"e2_{task_name}_{rep_name}_{js_name}_{norm}_taskavg"; count += 1
                result = evaluate_pca(data_array, label, jids=jids, norm=norm)
                if result and (result['pc2_var'] > 0.05 or result['rot_sum'] > 1.5):
                    print(f"  [{count}] {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                if result: results.append(result)
                
                # Option B: per-participant avg (each participant = 1 sample, averaged over all tasks)
                subj_avg_list = []
                for subj, task_dict in by_subj_task.items():
                    task_means = [np.mean(vecs, axis=0) for vecs in task_dict.values()]  # 15D per task
                    subj_avg_list.append(np.mean(task_means, axis=0))  # 15D per subject
                subj_avg = np.array(subj_avg_list)
                if len(subj_avg) >= 3:
                    label = f"e2_{task_name}_{rep_name}_{js_name}_{norm}_subjavg"; count += 1
                    result = evaluate_pca(subj_avg, label, jids=jids, norm=norm)
                    if result and (result['pc2_var'] > 0.05 or result['rot_sum'] > 1.5):
                        print(f"  [{count}] {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    if result: results.append(result)
    return results

def run_exp3():
    print(f"\n{'='*60}\nEXP 3: NORM DEEP DIVE\n{'='*60}")
    results = []
    task_sets = {'power_precision': sorted(set(POWER_IDS)|set(PRECISION_IDS)), 'all_grasp': ALL_GRASP_IDS, 'all_tasks': sorted(set(ALL_TASKS)-{0})}
    norms = ['none', 'zscore', 'range', 'robust', 'mean']
    for task_name, task_ids in task_sets.items():
        print(f"\n  {task_name}...", end='', flush=True)
        trials = load_myoki_trials(FILES, task_ids); print(f" {len(trials)} trials")
        for rep_name, rep_fn in [('trialmean', extract_trial_means), ('static5', lambda t: extract_static(t,5)), ('trialsd', extract_trial_std), ('trialrange', extract_trial_range)]:
            processed = rep_fn(trials); data_array = np.vstack([p[0] for p in processed])
            for norm in norms:
                for js_name, jids in [('all15', None), ('flex10', FLEXION_IDS), ('no_abd', [0,1,2,3,4,5,6,7,9,10,12,13])]:
                    label = f"e3_{task_name}_{rep_name}_{js_name}_{norm}"
                    result = evaluate_pca(data_array, label, jids=jids, norm=norm)
                    if result and (result['pc2_var'] > 0.05 or result['rot_sum'] > 1.5):
                        print(f"    {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    if result: results.append(result)
    return results

def load_ninapro_exercise(ex, hand_only=True):
    pre_dir = 'Ninapro Data/preprocessed'; n_j = 20 if hand_only else 22
    all_data = []; subject_data = {}
    files = sorted(glob.glob(os.path.join(pre_dir, f'S*_{ex}_preprocessed.mat')))
    for fpath in files:
        subj = os.path.basename(fpath).split('_')[0]
        mat = sio.loadmat(fpath); angles = mat['angles']
        if hand_only: angles = angles[:, :n_j]
        all_data.append(angles); subject_data[subj] = angles
    combined = np.vstack(all_data) if all_data else np.array([])
    return combined, subject_data

def run_exp4():
    print(f"\n{'='*60}\nEXP 4: NINAPRO\n{'='*60}")
    results = []
    ninapro = {}
    for ex in ['E1','E2','E3']:
        c, s = load_ninapro_exercise(ex); ninapro[ex] = (c, s)
        if len(c) > 0: print(f"  Ninapro {ex}: {len(c)} samples, {len(s)} subjects")
    for ex, (data, subjects) in ninapro.items():
        if len(data) == 0: continue
        for norm in ['none', 'zscore', 'range', 'robust']:
            label = f"e4_ninapro_{ex}_all_{norm}"
            result = evaluate_pca(data, label, norm=norm)
            if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
        for subj, sdata in subjects.items():
            for norm in ['none','zscore']:
                label = f"e4_ninapro_{subj}_{ex}_{norm}"
                result = evaluate_pca(sdata, label, norm=norm)
                if result:
                    if result['pc2_var'] > 0.08: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    results.append(result)
    combos = [('E12',['E1','E2']),('E13',['E1','E3']),('E23',['E2','E3']),('E123',['E1','E2','E3'])]
    for name, ex_list in combos:
        data_list = [ninapro[e][0] for e in ex_list if len(ninapro[e][0])>0]
        if data_list:
            combined = np.vstack(data_list)
            for norm in ['none','zscore','range']:
                label = f"e4_ninapro_{name}_{norm}"
                result = evaluate_pca(combined, label, norm=norm)
                if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
    return results

def run_exp6():
    print(f"\n{'='*60}\nEXP 6: PARTICIPANT SELECTION\n{'='*60}")
    results = []
    task_sets = [('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    participant_scores = {}
    for tn, ids in task_sets:
        for fpath in FILES:
            subj = os.path.basename(fpath).replace('.mat','')
            trials = load_myoki_trials([fpath], ids)
            if len(trials) < 5: continue
            for norm in ['none','zscore']:
                data = np.vstack([np.mean(t[0],axis=0) for t in trials])
                label = f"e6_{subj}_{tn}_trialmean_{norm}"
                result = evaluate_pca(data, label, norm=norm)
                if result:
                    results.append(result)
                    if result['pc2_var'] > 0.07: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    if subj not in participant_scores: participant_scores[subj] = []
                    participant_scores[subj].append(result['pc2_var'])
    if participant_scores:
        best_pc2 = sorted(participant_scores.items(), key=lambda x: np.mean(x[1]), reverse=True)
        print(f"\n  Best by PC2:")
        for s, scores in best_pc2[:5]: print(f"    {s}: avg PC2={np.mean(scores)*100:.1f}%")
        top3 = [s for s,_ in best_pc2[:3]]
        for tn, ids in task_sets:
            top_files = [f for f in FILES if any(s in os.path.basename(f) for s in top3)]
            trials = load_myoki_trials(top_files, ids)
            for norm in ['none','zscore','range']:
                data = np.vstack([np.mean(t[0],axis=0) for t in trials])
                label = f"e6_top3pc2_{tn}_trialmean_{norm}"
                result = evaluate_pca(data, label, norm=norm)
                if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
                processed = extract_static(trials,5); data = np.vstack([p[0] for p in processed])
                label = f"e6_top3pc2_{tn}_static5_{norm}"
                result = evaluate_pca(data, label, norm=norm)
                if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
    return results

def run_exp7():
    print(f"\n{'='*60}\nEXP 7: STATIC+DYNAMIC FEATURES\n{'='*60}")
    results = []
    task_sets = [('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    for tn, ids in task_sets:
        print(f"\n  {tn}...", end='', flush=True)
        trials = load_myoki_trials(FILES, ids); print(f" {len(trials)} trials")
        trialmeans = np.vstack([np.mean(t[0],axis=0) for t in trials])
        trialstd = np.vstack([np.std(t[0],axis=0) for t in trials])
        trialrange = np.vstack([np.max(t[0],axis=0)-np.min(t[0],axis=0) for t in trials])
        for norm in ['none','zscore','range']:
            data = np.hstack([trialmeans, trialstd])
            label = f"e7_{tn}_mean_std_{norm}"
            result = evaluate_pca(data, label, norm=norm)
            if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
            data = np.hstack([trialmeans, trialrange])
            label = f"e7_{tn}_mean_range_{norm}"
            result = evaluate_pca(data, label, norm=norm)
            if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
        for norm2 in ['none','zscore']:
            label = f"e7_{tn}_stdonly_{norm2}"
            result = evaluate_pca(trialstd, label, norm=norm2)
            if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
    return results

def run_exp8():
    print(f"\n{'='*60}\nEXP 8: TASK BALANCING\n{'='*60}")
    results = []
    n_power = len(POWER_IDS); n_prec = len(PRECISION_IDS)
    mixes = [('pw25_pr75', POWER_IDS[:max(1,n_power//4)], PRECISION_IDS[:max(1,3*n_prec//4)]),
             ('pw50_pr50', POWER_IDS[:max(1,n_power//2)], PRECISION_IDS[:max(1,n_prec//2)]),
             ('pw75_pr25', POWER_IDS[:max(1,3*n_power//4)], PRECISION_IDS[:max(1,n_prec//4)]),
             ('pwEq_prEq', POWER_IDS[:min(n_power,n_prec)], PRECISION_IDS[:min(n_power,n_prec)]),
             ('pwAll_pr10', POWER_IDS, PRECISION_IDS[:10]),
             ('pw10_prAll', POWER_IDS[:10], PRECISION_IDS),
             ('pw5_pr5', POWER_IDS[:5], PRECISION_IDS[:5])]
    for mix_name, power_ids, prec_ids in mixes:
        combined_ids = sorted(set(power_ids)|set(prec_ids))
        print(f"\n  {mix_name}: {len(power_ids)}p+{len(prec_ids)}r={len(combined_ids)} tasks...", end='', flush=True)
        trials = load_myoki_trials(FILES, combined_ids); print(f" {len(trials)} trials")
        if len(trials) < 5: continue
        for rep_name, rep_fn in [('trialmean', extract_trial_means), ('static5', lambda t: extract_static(t,5)), ('trialsd', extract_trial_std)]:
            processed = rep_fn(trials); data = np.vstack([p[0] for p in processed])
            for norm in ['none','zscore','range']:
                for js_name, jids in [('all15',None),('flex10',FLEXION_IDS),('no_abd',[0,1,2,3,4,5,6,7,9,10,12,13])]:
                    label = f"e8_{mix_name}_{rep_name}_{js_name}_{norm}"
                    result = evaluate_pca(data, label, jids=jids, norm=norm)
                    if result is None: continue
                    if result['pc2_var'] > 0.06 or result['rot_sum'] > 1.5:
                        print(f"    {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    results.append(result)
    return results

def run_exp9():
    print(f"\n{'='*60}\nEXP 9: ONE-SAMPLE-PER-TASK\n{'='*60}")
    results = []
    task_sets = [('power_only', POWER_IDS), ('precision_only', PRECISION_IDS), ('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    for tn, ids in task_sets:
        print(f"\n  {tn} ({len(ids)} tasks)...", end='', flush=True)
        trials = load_myoki_trials(FILES, ids); print(f" {len(trials)} trials")
        by_task = {}
        for t, tid, subj in trials:
            if tid not in by_task: by_task[tid] = []
            by_task[tid].append(t)
        print(f"    {len(by_task)} tasks with data")
        for rep_name, rep_fn in [('trialmean', lambda v: np.mean(v,axis=0)), ('static5', lambda v: np.mean(v[-5:],axis=0)), ('trialsd', lambda v: np.std(v,axis=0)), ('trialrange', lambda v: np.max(v,axis=0)-np.min(v,axis=0))]:
            task_data = [np.mean([rep_fn(t) for t in trial_list], axis=0) for tid, trial_list in by_task.items()]
            data_array = np.array(task_data)
            for norm in ['none','zscore','range','robust']:
                for js_name, jids in [('all15',None),('flex10',FLEXION_IDS),('no_abd',[0,1,2,3,4,5,6,7,9,10,12,13])]:
                    label = f"e9_{tn}_{rep_name}_{js_name}_{norm}"
                    result = evaluate_pca(data_array, label, jids=jids, norm=norm)
                    if result is None: continue
                    if result['pc2_var'] > 0.06 or result['rot_sum'] > 1.5:
                        print(f"    {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                    results.append(result)
    return results

def run_exp10():
    print(f"\n{'='*60}\nEXP 10: REST POSTURE SUBTRACTION\n{'='*60}")
    results = []
    rest_poses = []
    for fpath in FILES:
        mat = sio.loadmat(fpath); glove = np.nan_to_num(mat['glove_calibrated'])
        tasks = mat['task'].flatten().astype(int)
        rest = glove[tasks == 0, :15]
        if len(rest) > 0: rest_poses.append(np.mean(rest, axis=0))
    global_rest = np.mean(rest_poses, axis=0) if rest_poses else None
    print(f"  Rest posture: {global_rest}")
    task_sets = [('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    for tn, ids in task_sets:
        trials = load_myoki_trials(FILES, ids)
        trialmeans = np.vstack([np.mean(t[0], axis=0) for t in trials])
        if global_rest is not None:
            rest_sub = trialmeans - global_rest
            for norm in ['none','zscore','range']:
                label = f"e10_{tn}_restsub_trialmean_{norm}"
                result = evaluate_pca(rest_sub, label, norm=norm)
                if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
        try:
            log_data = np.log(np.abs(trialmeans - global_rest) + 1) * np.sign(trialmeans - global_rest)
            for norm in ['none','zscore']:
                label = f"e10_{tn}_log_{norm}"
                result = evaluate_pca(log_data, label, norm=norm)
                if result: print(f"  {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
        except: pass
    return results

def run_exp11():
    print(f"\n{'='*60}\nEXP 11: ALL FRAMES\n{'='*60}")
    results = []
    task_sets = [('all_tasks', sorted(set(ALL_TASKS)-{0})), ('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    for tn, ids in task_sets:
        print(f"\n  {tn}...", end='', flush=True)
        trials = load_myoki_trials(FILES, ids); print(f" {len(trials)} trials")
        all_data = np.vstack([t[0] for t in trials])
        print(f"    {len(all_data)} frames")
        if len(all_data) > 50000:
            np.random.seed(42); idx = np.random.choice(len(all_data), 50000, replace=False); all_data = all_data[idx]
        for norm in ['none','zscore','range']:
            label = f"e11_{tn}_allframes_{norm}"
            result = evaluate_pca(all_data, label, norm=norm)
            if result: print(f"    {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}"); results.append(result)
    return results

def run_exp12():
    print(f"\n{'='*60}\nEXP 12: DISCRIMINATIVE JOINTS\n{'='*60}")
    results = []
    custom_subsets = {'mcp_only': [1,4,6,9,12], 'pip_only': [2,5,7,10,13], 'mcp_pip_no_thumb': [4,5,6,7,9,10,12,13],
                      'thumb_vs_little': [1,2,12,13], 'thumb_idx_lit': [1,2,4,5,12,13],
                      'flex_plus_opp': [1,2,4,5,6,7,9,10,12,13,3], 'thumb_opp_flex': [3,1,2,4,5,6,7,9,10,12,13]}
    task_sets = [('power_precision', sorted(set(POWER_IDS)|set(PRECISION_IDS))), ('all_grasp', ALL_GRASP_IDS)]
    for tn, ids in task_sets:
        print(f"\n  {tn}...", end='', flush=True)
        trials = load_myoki_trials(FILES, ids); print(f" {len(trials)} trials")
        trialmeans = np.vstack([np.mean(t[0],axis=0) for t in trials])
        for subset_name, jids in custom_subsets.items():
            for norm in ['none','zscore','range']:
                label = f"e12_{tn}_{subset_name}_{norm}"
                result = evaluate_pca(trialmeans, label, jids=jids, norm=norm)
                if result is None: continue
                if result['pc2_var'] > 0.08 or result['rot_sum'] > 1.4:
                    print(f"    {label:<50} V={result['pc1_var']*100:>4.1f}+{result['pc2_var']*100:>4.1f}={result['pc1_pc2_sum']*100:>4.1f}% S1={result['rot_s1_sim']:.4f} S2={result['rot_s2_sim']:.4f}")
                results.append(result)
    return results

def score_balanced(result):
    s1 = result.get('rot_s1_sim', 0); s2 = result.get('rot_s2_sim', 0)
    pc2 = result.get('pc2_var', 0); pc1 = result.get('pc1_var', 0); total = result.get('pc1_pc2_sum', 0)
    pattern_score = s1 + s2
    pc2_bonus = (pc2 - 0.08) * 5 if pc2 >= 0.08 else 0
    total_bonus = (total - 0.80) * 2 if total >= 0.80 else 0
    pc1_penalty = (pc1 - 0.90) * 5 if pc1 > 0.90 else 0
    return pattern_score + pc2_bonus + total_bonus - pc1_penalty

def print_summary_table(all_results, top_n=50):
    if not all_results: print("No results!"); return
    seen = set(); unique = []
    for r in all_results:
        if r['label'] not in seen: seen.add(r['label']); unique.append(r)
    print(f"\n{'='*130}")
    print(f"RANKING: {len(unique)} unique experiments (by balanced score)")
    print(f"{'='*130}")
    unique.sort(key=score_balanced, reverse=True)
    print(f"{'Rank':<5} {'Experiment':<55} {'PC1%':>7} {'PC2%':>7} {'Sum%':>7} {'S1':>8} {'S2':>8} {'BalSc':>7} {'N':>6} {'J':>3}")
    print("-" * 130)
    for i, r in enumerate(unique[:top_n]):
        bs = score_balanced(r); star = "★" if i == 0 else " "; fl = "!" if r['pc2_var']>0.08 else " "
        print(f"{star}{i+1:<4}{fl} {r['label']:<55} {r['pc1_var']*100:>6.1f}% {r['pc2_var']*100:>6.1f}% {r['pc1_pc2_sum']*100:>6.1f}% {r['rot_s1_sim']:>+7.4f} {r['rot_s2_sim']:>+7.4f} {bs:>6.3f} {r['n_samples']:>6} {r['n_joints']:>3}")
    clean = []
    for r in unique:
        c = {}
        for k, v in r.items():
            if isinstance(v, (np.floating,)): c[k] = float(v)
            elif isinstance(v, (np.integer,)): c[k] = int(v)
            elif isinstance(v, np.ndarray): c[k] = v.tolist()
            else: c[k] = v
        c['balance_score'] = float(score_balanced(r))
        clean.append(c)
    save_path = os.path.join(SAVE_DIR, 'mega_balance_results.json')
    with open(save_path, 'w') as f: json.dump(clean, f, indent=2)
    print(f"\nSaved {len(clean)} results to {save_path}")
    print(f"\nTOP 10 BY PATTERN SIMILARITY:")
    by_pattern = sorted(unique, key=lambda r: r['rot_s1_sim']+r['rot_s2_sim'], reverse=True)
    for i, r in enumerate(by_pattern[:10]): print(f"  {i+1}. {r['label']:<55} S1={r['rot_s1_sim']:.4f} S2={r['rot_s2_sim']:.4f} V={r['pc1_var']*100:.1f}+{r['pc2_var']*100:.1f}")
    print(f"\nTOP 10 BY PC2 VARIANCE:")
    by_pc2 = sorted(unique, key=lambda r: r['pc2_var'], reverse=True)
    for i, r in enumerate(by_pc2[:10]): print(f"  {i+1}. {r['label']:<55} PC2={r['pc2_var']*100:.1f}% S1={r['rot_s1_sim']:.4f} S2={r['rot_s2_sim']:.4f}")
    print(f"\nBEST BALANCED (PC2>8%, S1>0.7):")
    balanced = [r for r in unique if r['pc2_var']>0.08 and r['rot_s1_sim']>0.70]
    balanced.sort(key=lambda r: r['rot_s1_sim']+r['rot_s2_sim']+r['pc2_var']*5, reverse=True)
    for i, r in enumerate(balanced[:15]): print(f"  {i+1}. {r['label']:<55} V={r['pc1_var']*100:.1f}+{r['pc2_var']*100:.1f}={r['pc1_pc2_sum']*100:.1f} S1={r['rot_s1_sim']:.4f} S2={r['rot_s2_sim']:.4f}")
    return unique

def save_best_results(unique_results):
    unique_results.sort(key=score_balanced, reverse=True)
    good = [r for r in unique_results if r['rot_s1_sim']>0.65 and r['pc1_pc2_sum']>0.70]
    by_pc2 = sorted(unique_results, key=lambda r: r['pc2_var'], reverse=True)
    good.extend([r for r in by_pc2[:5] if r not in good and r['pc2_var']>0.08])
    good = good[:20]
    print(f"\n{'='*60}\nSAVING {len(good)} BEST AS NPZ\n{'='*60}")
    saved = []
    for idx, r in enumerate(good):
        n_j = r['n_joints']; pc1 = np.array(r['pc1']); pc2 = np.array(r['pc2'])
        if 'jids' in r and r['jids'] is not None:
            jids = [j for j in r['jids'] if j < 15]  # Only keep valid indices for 15-joint output
            full_components = np.zeros((2, 15))
            if jids:
                k = min(len(jids), len(pc1))
                full_components[0, jids[:k]] = pc1[:k]
                if len(pc2) >= k: full_components[1, jids[:k]] = pc2[:k]
        else:
            full_components = np.zeros((2, 15))
            k = min(n_j, 15, len(pc1))
            full_components[0, :k] = pc1[:k]
            if len(pc2) >= k: full_components[1, :k] = pc2[:k]

        save_path = os.path.join('PCA_results/experiments', f"balance_{idx+1:02d}.npz")
        rot_s1 = np.array(r['rot_s1']); rot_s2 = np.array(r['rot_s2'])
        np.savez(save_path, components=full_components, rot_s1=rot_s1, rot_s2=rot_s2,
                 pc1=pc1, pc2=pc2, explained_variance=np.array([r['pc1_var'], r['pc2_var']]),
                 rot_s1_sim=r['rot_s1_sim'], rot_s2_sim=r['rot_s2_sim'],
                 pc1_var=r['pc1_var'], pc2_var=r['pc2_var'], total_var=r['pc1_pc2_sum'],
                 label=r['label'], n_joints=n_j)
        print(f"  {idx+1}. {r['label'][:55]:<57} -> balance_{idx+1:02d}.npz")
        saved.append({'file': f"balance_{idx+1:02d}.npz", 'label': r['label'],
                      'var_pc1': r['pc1_var'], 'var_pc2': r['pc2_var'], 'var_sum': r['pc1_pc2_sum'],
                      's1_sim': r['rot_s1_sim'], 's2_sim': r['rot_s2_sim'], 'balance_score': score_balanced(r)})
    with open(os.path.join(SAVE_DIR, 'balance_manifest.json'), 'w') as f: json.dump(saved, f, indent=2)
    print(f"Manifest saved")

if __name__ == '__main__':
    print("="*60); print("MEGA BALANCE SEARCH"); print(f"Started: {datetime.now():%Y-%m-%d %H:%M:%S}"); print("="*60)
    all_results = []
    all_results.extend(run_exp1())
    all_results.extend(run_exp2())
    all_results.extend(run_exp3())
    all_results.extend(run_exp4())
    all_results.extend(run_exp6())
    all_results.extend(run_exp7())
    all_results.extend(run_exp8())
    all_results.extend(run_exp9())
    all_results.extend(run_exp10())
    all_results.extend(run_exp11())
    all_results.extend(run_exp12())
    print(f"\n{'='*60}\nTOTAL: {len(all_results)} results\n{'='*60}")
    unique = print_summary_table(all_results, top_n=60)
    save_best_results(unique)
    print(f"\nDone! {datetime.now():%Y-%m-%d %H:%M:%S}")
